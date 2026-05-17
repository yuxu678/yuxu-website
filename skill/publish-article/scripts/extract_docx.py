#!/usr/bin/env python3
"""Extract article content from a Word (.docx) file and output as JSON.

Usage:
    python extract_docx.py <input.docx> [--date YYYY-MM-DD] [--title "Custom Title"]

Output: JSON with {id, date, title, content} to stdout.

Dependencies: python-docx (pip install python-docx)
"""

import argparse
import json
import sys
from datetime import date
from pathlib import Path

try:
    from docx import Document
except ImportError:
    print("Please install python-docx: pip install python-docx", file=sys.stderr)
    sys.exit(1)


def extract_title(doc, doc_path, args_title=None):
    if args_title:
        return args_title
    for p in doc.paragraphs:
        style_name = p.style.name if p.style else ""
        if style_name.startswith("Heading"):
            title = p.text.strip()
            if title:
                return title
    return Path(doc_path).stem


def para_to_html(para):
    text = para.text.strip()
    if not text:
        return ""
    style_name = para.style.name if para.style else ""

    if style_name.startswith("Heading 1"):
        return "<h2>" + text + "</h2>"
    elif style_name.startswith("Heading 2"):
        return "<h3>" + text + "</h3>"
    elif style_name.startswith("Heading 3"):
        return "<h4>" + text + "</h4>"

    if "List" in style_name:
        return "<li>" + text + "</li>"

    has_bold = any(run.bold for run in para.runs if run.text.strip())
    if has_bold:
        parts = []
        for run in para.runs:
            if run.bold and run.text.strip():
                parts.append("<strong>" + run.text + "</strong>")
            else:
                parts.append(run.text)
        return "<p>" + "".join(parts) + "</p>"

    return "<p>" + text + "</p>"


def extract_content(doc):
    html_parts = []
    in_list = False

    for para in doc.paragraphs:
        html = para_to_html(para)
        if not html:
            continue

        if html.startswith("<li>"):
            if not in_list:
                html_parts.append("<ul>")
                in_list = True
            html_parts.append(html)
        else:
            if in_list:
                html_parts.append("</ul>")
                in_list = False
            html_parts.append(html)

    if in_list:
        html_parts.append("</ul>")

    for table in doc.tables:
        html_parts.append('<figure class="wp-block-table"><table>')
        for row in table.rows:
            html_parts.append("<tr>")
            for cell in row.cells:
                html_parts.append("<td>" + cell.text + "</td>")
            html_parts.append("</tr>")
        html_parts.append("</table></figure>")

    return "\n\n".join(html_parts)


def main():
    parser = argparse.ArgumentParser(description="Extract article from Word document")
    parser.add_argument("input", help="Path to .docx file")
    parser.add_argument("--date", default=str(date.today()), help="Article date (YYYY-MM-DD)")
    parser.add_argument("--title", help="Custom title (default: first heading or filename)")
    args = parser.parse_args()

    doc = Document(args.input)
    title = extract_title(doc, args.input, args.title)
    content_html = extract_content(doc)
    article_id = str(int(date.today().strftime("%Y%m%d")))

    article = {
        "id": article_id,
        "date": args.date,
        "title": title,
        "content": content_html,
    }

    print(json.dumps(article, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()